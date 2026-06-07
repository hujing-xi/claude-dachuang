"""RAG项目PPT v2 — 正确幻灯片顺序 + 校徽 + matplotlib配图"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

# ── 从模板加载 layouts ───────────────────────────────────────
TPL = "/root/.claude/uploads/11babfda-606b-5c7b-b135-f3a1b0077dbb/9a08159b-lll.pptx"
tpl = Presentation(TPL)
section_layout  = next(l for l in tpl.slide_layouts if l.name == "1_自定义版式")
blank_layout    = next(l for l in tpl.slide_layouts if l.name == "空白")
content_layout  = next(l for l in tpl.slide_layouts if l.name == "仅标题")

# 以原模板为基础，删除所有已有 slides
prs = tpl
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    rId = sldId.get(qn('r:id'))
    try: prs.part.drop_rel(rId)
    except: pass
    sldIdLst.remove(sldId)

# ── Colors ───────────────────────────────────────────────────
DEEP_BLUE  = RGBColor(0x07,0x1F,0x65);  C1 = DEEP_BLUE
TEAL       = RGBColor(0x34,0x8B,0x9F);  C2 = TEAL
RED        = RGBColor(0xC0,0x00,0x00);  C3 = RED
WHITE      = RGBColor(0xFF,0xFF,0xFF)
GRAY_BG    = RGBColor(0xE0,0xE0,0xE0)
LIGHT_GRAY = RGBColor(0xF2,0xF2,0xF2)
MID_GRAY   = RGBColor(0xAA,0xAA,0xAA)
DARK_TEXT  = RGBColor(0x22,0x22,0x22)
ROW_ALT    = RGBColor(0xEB,0xF3,0xFB)
DARK_SIDE  = RGBColor(0x44,0x44,0x44)

W = prs.slide_width
H = prs.slide_height

# ── 图片路径 ─────────────────────────────────────────────────
EMB_CONTENT = "/tmp/pptx_imgs/slide4_图片_18_94bc48a2.jpg"
EMB_SECTION = "/tmp/pptx_imgs/slide7_图片_12_48f5857e.jpg"
COVER_LOGO  = "/tmp/pptx_imgs/slide1_图片_9_b589cbf7.jpg"
D = {1:"/tmp/diagrams/d1_timeline.png", 2:"/tmp/diagrams/d2_landscape.png",
     3:"/tmp/diagrams/d3_arch.png",     4:"/tmp/diagrams/d4_router.png",
     5:"/tmp/diagrams/d5_gantt.png",    6:"/tmp/diagrams/d6_knowledge.png",
     7:"/tmp/diagrams/d7_retrieval.png"}

# ── 基础绘图函数 ─────────────────────────────────────────────
def add(layout): return prs.slides.add_slide(layout)

def rect(sl, x, y, w, h, fc=None, ec=None, lw=None):
    s = sl.shapes.add_shape(1, x, y, w, h)
    if fc:
        s.fill.solid(); s.fill.fore_color.rgb = fc
    else:
        s.fill.background()
    if ec:
        s.line.color.rgb = ec
        if lw: s.line.width = lw
    else:
        s.line.fill.background()
    return s

def tb(sl, x, y, w, h, text, sz=18, bold=False, color=DARK_TEXT,
       align=PP_ALIGN.LEFT, wrap=True):
    box = sl.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame; tf.word_wrap = wrap
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name = "微软雅黑"; r.font.size = Pt(sz)
    r.font.bold = bold; r.font.color.rgb = color
    return box

def pic(sl, path, x, y, w, h=None):
    if h: return sl.shapes.add_picture(path, x, y, w, h)
    return sl.shapes.add_picture(path, x, y, w)

def oval(sl, x, y, w, h, fc):
    s = sl.shapes.add_shape(9, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fc; s.line.fill.background()
    return s

def emblem(sl):
    try: pic(sl, EMB_CONTENT, Inches(0.19), Inches(0.09), Inches(0.84), Inches(0.84))
    except: pass

def header(sl, title, page_num):
    rect(sl, Inches(1.0), Inches(0.1), Inches(0.12), Inches(0.82), fc=TEAL)
    rect(sl, Inches(1.12), Inches(0.1), Inches(12.0), Inches(0.82), fc=DEEP_BLUE)
    tb(sl, Inches(1.28), Inches(0.16), Inches(11.7), Inches(0.7),
       title, sz=26, bold=True, color=WHITE)
    tb(sl, Inches(11.8), Inches(7.0), Inches(1.3), Inches(0.4),
       str(page_num), sz=15, color=MID_GRAY, align=PP_ALIGN.RIGHT)
    emblem(sl)

def refline(sl, text):
    tb(sl, Inches(0.4), Inches(7.1), Inches(12.5), Inches(0.3),
       text, sz=11, color=MID_GRAY)

def section_page(part_n, chapter_name):
    s = add(section_layout)
    tb(s, Inches(1.2), Inches(1.8), Inches(3.8), Inches(1.1),
       f"PART  0{part_n}", sz=28, bold=True, color=DEEP_BLUE)
    rect(s, Inches(1.2), Inches(3.1), Inches(3.8), Inches(0.07), fc=TEAL)
    tb(s, Inches(5.5), Inches(1.45), Inches(2.8), Inches(1.1),
       f"0{part_n}", sz=72, bold=True, color=WHITE)
    tb(s, Inches(5.5), Inches(2.55), Inches(7.6), Inches(0.8),
       f"Part {part_n}", sz=20, bold=True, color=WHITE)
    tb(s, Inches(8.0), Inches(2.85), Inches(5.1), Inches(1.2),
       chapter_name, sz=38, bold=True, color=WHITE)
    try: pic(s, EMB_SECTION, Inches(1.85), Inches(1.91), Inches(1.9), Inches(1.9))
    except: pass

# ═══════════════════════════════════════════════════════════════
# SLIDE 1 — COVER
# ═══════════════════════════════════════════════════════════════
s = add(blank_layout)
rect(s, 0, 0, W, H, fc=GRAY_BG)
rect(s, 0, 0, Inches(0.5), H, fc=DARK_SIDE)
try: pic(s, COVER_LOGO, Inches(0.58), Inches(0.12), Inches(3.6), Inches(1.53))
except: pass
rect(s, Inches(0.7), Inches(2.05), Inches(12.3), Inches(5.0), fc=WHITE, ec=RGBColor(0xCC,0xCC,0xCC))
rect(s, Inches(0.7), Inches(2.55), Inches(12.3), Inches(1.48), fc=DEEP_BLUE)
tb(s, Inches(0.9), Inches(2.62), Inches(12.0), Inches(1.35),
   "面向垂直场景的智能体RAG知识服务系统\n研究与实现",
   sz=32, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
rect(s, Inches(0.7), Inches(4.03), Inches(12.3), Inches(0.07), fc=RED)
tb(s, Inches(0.7), Inches(4.18), Inches(12.3), Inches(0.6),
   "���学生创新创业训练计划立项汇报",
   sz=22, color=DEEP_BLUE, align=PP_ALIGN.CENTER)
rect(s, Inches(2.5), Inches(4.93), Inches(8.0), Inches(0.03), fc=MID_GRAY)
tb(s, Inches(0.7), Inches(5.08), Inches(12.3), Inches(0.55),
   "���责人：胡景熙   |   成员：孙泽恺、冷诗雨   |   指导教师：周航",
   sz=18, color=DARK_TEXT, align=PP_ALIGN.CENTER)
tb(s, Inches(0.7), Inches(5.68), Inches(12.3), Inches(0.55),
   "北京交通大学威海国际学院   |   信息管理与信息系统   |   2026年6月",
   sz=17, color=MID_GRAY, align=PP_ALIGN.CENTER)
rect(s, Inches(0.7), Inches(6.65), Inches(12.3), Inches(0.3), fc=TEAL)
tb(s, Inches(0.9), Inches(6.68), Inches(12.0), Inches(0.26),
   "创新训练项目  ·  成果形式：学术论文 + 软件程序",
   sz=13, color=WHITE, align=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════════════════
# SLIDE 2 — TOC
# ═══════════════════════════════════════════════════════════════
s = add(blank_layout)
rect(s, 0, 0, W, H, fc=WHITE)
rect(s, 0, 0, Inches(0.5), H, fc=DEEP_BLUE)
rect(s, Inches(0.5), Inches(0.2), Inches(0.12), Inches(0.85), fc=RED)
rect(s, Inches(0.62), Inches(0.2), Inches(4.5), Inches(0.85), fc=DEEP_BLUE)
tb(s, Inches(0.78), Inches(0.26), Inches(4.3), Inches(0.72),
   "目录  CONTENTS", sz=26, bold=True, color=WHITE)
rect(s, Inches(12.0), Inches(0.2), Inches(1.0), Inches(7.1), fc=DEEP_BLUE)
rect(s, Inches(12.4), Inches(0.2), Inches(0.6), Inches(7.1), fc=TEAL)

toc = [("01","选题背景与意义","研究背景、选题意义与延续性",DEEP_BLUE),
       ("02","研究现状分析","国内外进展与现有方法局限",TEAL),
       ("03","主要研究内容","系统架构、路由层、多策略检索与验证",RED),
       ("04","工作计划与预期成果","项目进度安排、预期论文与软件成果",DEEP_BLUE),
       ("05","参考文献","12篇核心文献",TEAL)]
for i,(num,label,sub,col) in enumerate(toc):
    y = Inches(1.38 + i * 1.1)
    rect(s, Inches(0.65), y, Inches(11.2), Inches(0.95), fc=LIGHT_GRAY)
    rect(s, Inches(0.65), y, Inches(0.08), Inches(0.95), fc=col)
    oval(s, Inches(0.85), y+Inches(0.18), Inches(0.6), Inches(0.6), col)
    tb(s, Inches(0.82), y+Inches(0.14), Inches(0.65), Inches(0.58),
       num, sz=17, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tb(s, Inches(1.62), y+Inches(0.08), Inches(6.0), Inches(0.45),
       label, sz=21, bold=True, color=col)
    tb(s, Inches(1.62), y+Inches(0.54), Inches(9.5), Inches(0.36),
       sub, sz=13, color=MID_GRAY)

# ═══════════════════════════════════════════════════════════════
# SLIDE 3 — Part 1 过渡
# ═══════════════════════════════════════════════════════════════
section_page(1, "选题背景与意义")

# ═══════════════════════════════════════════════════════════════
# SLIDE 4 — 1-1 选题背景
# ═══════════════════════════════════════════════════════════════
s = add(content_layout)
header(s, "1-1  选题背景", 4)
pic(s, D[1], Inches(0.4), Inches(1.08), Inches(12.5), Inches(3.2))
for i,(col,ttl,body) in enumerate([
    (C1,"LLM 快速发展，RAG 成为主流",
     "大语言模型能力迅速提升，检索增强生成（RAG）已成为知识密集型问答主流技术路线 [1][10]"),
    (TEAL,"垂直场景面临异质性挑战",
     "景区/校园等场景同时存在事实查询、空间推理、路线规划三类任务，单一RAG路径无法应对")]):
    x = Inches(0.4 + i * 6.45)
    rect(s, x, Inches(4.45), Inches(6.1), Inches(2.75), fc=LIGHT_GRAY)
    rect(s, x, Inches(4.45), Inches(0.1), Inches(2.75), fc=col)
    tb(s, x+Inches(0.18), Inches(4.58), Inches(5.8), Inches(0.5),
       ttl, sz=17, bold=True, color=col)
    tb(s, x+Inches(0.18), Inches(5.15), Inches(5.8), Inches(1.95),
       body, sz=14, color=DARK_TEXT)
refline(s, "引用：[1] Lewis et al. NeurIPS 2020  [6] Dong et al. ACL 2025  [10] Gao et al. 2024")

# ═══════════════════════════════════════════════════════════════
# SLIDE 5 — 1-2 选题意义
# ═══════════════════════════════════════════════════════════════
s = add(content_layout)
header(s, "1-2  选题意义", 5)
for i,(col,ttl,body) in enumerate([
    (DEEP_BLUE,"理论意义",
     "首次针对景区服务场景中用户请求异质性进行系统建模，将请求划分为三类并设计对应检索路径，填补多策略RAG研究空白。"),
    (TEAL,"应用价值",
     "为洱海生态廊道等景区提供可复用的智能知识服务框架，支持事实问答与空间关系推理，具有直接落地意义。"),
    (RED,"研究延续性",
     "在团队去年大创事实问答模块基础上系统扩展：引入图谱检索、多策略路由与验证层，研究脉络完整。")]):
    cx = Inches(0.35 + i * 4.32)
    cw = Inches(4.0)
    rect(s, cx, Inches(1.12), cw, Inches(6.08), fc=LIGHT_GRAY)
    rect(s, cx, Inches(1.12), cw, Inches(0.08), fc=col)
    oval(s, cx+Inches(1.5), Inches(1.38), Inches(1.0), Inches(1.0), col)
    tb(s, cx+Inches(1.45), Inches(1.4), Inches(1.0), Inches(0.96),
       str(i+1), sz=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    rect(s, cx+Inches(0.4), Inches(2.53), cw-Inches(0.8), Inches(0.05), fc=col)
    tb(s, cx+Inches(0.12), Inches(2.62), cw-Inches(0.24), Inches(0.6),
       ttl, sz=21, bold=True, color=col, align=PP_ALIGN.CENTER)
    tb(s, cx+Inches(0.15), Inches(3.32), cw-Inches(0.3), Inches(3.75),
       body, sz=15, color=DARK_TEXT)

# ═══════════════════════════════════════════════════════════════
# SLIDE 6 — Part 2 过渡
# ═══════════════════════════════════════════════════════════════
section_page(2, "研究现状分析")

# ═══════════════════════════════════════════════════════════════
# SLIDE 7 — 2-1 国内外研究现状
# ═══════════════════════════════════════════════════════════════
s = add(content_layout)
header(s, "2-1  国内外研究现状", 7)
pic(s, D[2], Inches(0.4), Inches(1.05), Inches(12.5), Inches(5.9))
refline(s, "引用：[1][8][9][10][11]")

# ═══════════════════════════════════════════════════════════════
# SLIDE 8 — 2-2 现有方法三大不足
# ═══════════════════════════════════════════════════════════════
s = add(content_layout)
header(s, "2-2  现有方法的三大不足", 8)
for i,(col,ttl,body) in enumerate([
    (RED,"❶  任务类型无法区分",
     "现有RAG对事实查询、空间推理、路线规划采用统一策略，导致非事实类任务准确率大幅下降。"),
    (DEEP_BLUE,"❷  检索策略缺乏差异化",
     "单一向量检索对图结构关系问题效果有限；纯关键词检索对语义相似查询召回率低，两端均不适用。"),
    (TEAL,"❸  生成结果缺乏可靠验证",
     "大多数RAG缺少事后验证，面对时效冲突或约束违反时无法主动纠错，幻觉风险高 [3][4][11]。")]):
    gy = Inches(1.25 + i * 1.6)
    rect(s, Inches(0.4), gy, Inches(12.5), Inches(1.45), fc=LIGHT_GRAY)
    rect(s, Inches(0.4), gy, Inches(0.12), Inches(1.45), fc=col)
    tb(s, Inches(0.62), gy+Inches(0.1), Inches(12.0), Inches(0.5),
       ttl, sz=20, bold=True, color=col)
    tb(s, Inches(0.62), gy+Inches(0.62), Inches(12.0), Inches(0.75),
       body, sz=15, color=DARK_TEXT)
rect(s, Inches(0.4), Inches(6.3), Inches(12.5), Inches(0.82), fc=DEEP_BLUE)
tb(s, Inches(0.6), Inches(6.38), Inches(12.2), Inches(0.65),
   "→  本项目提出「多策略 Agentic RAG 框架」，针对性解决以上三大问题，填补��直场景研究空白",
   sz=18, bold=True, color=WHITE)
refline(s, "引用：[2][3][5][6][7][9]")

# ═══════════════════════════════════════════════════════════════
# SLIDE 9 — Part 3 过渡
# ═══════════════════════════════════════════════════════════════
section_page(3, "主要研究内容")

# ═══════════════════════════════════════════════════════════════
# SLIDE 10 — 3-0 系统架构总览
# ═══════════════════════════════════════════════════════════════
s = add(content_layout)
header(s, "3-0  系统架构总览", 10)
pic(s, D[3], Inches(0.4), Inches(1.05), Inches(12.5), Inches(5.9))

# ═══════════════════════════════════════════════════════════════
# SLIDE 11 — 3-1 知识底座构建
# ═══════════════════════════════════════════════════════════════
s = add(content_layout)
header(s, "3-1  场景化知识底座构建", 11)
rect(s, Inches(0.4), Inches(1.1), Inches(4.5), Inches(0.55), fc=RED)
tb(s, Inches(0.5), Inches(1.12), Inches(4.3), Inches(0.5),
   "针对问题", sz=19, bold=True, color=WHITE)
rect(s, Inches(0.4), Inches(1.65), Inches(4.5), Inches(5.05), fc=LIGHT_GRAY)
tb(s, Inches(0.55), Inches(1.78), Inches(4.2), Inches(4.8),
   "景区数据散乱异构，缺乏统一知识表示：\n\n"
   "• 文本资料（非结构化）\n  介绍手册、规则公告文本\n\n"
   "• 空间数据（半结构化）\n  节点坐标、设施属性信息\n\n"
   "• 关系数据（结构化）\n  景点关联、路径约束条件\n\n"
   "→ 三类格式差异导致统一检索困难",
   sz=14, color=DARK_TEXT)
tb(s, Inches(5.1), Inches(3.5), Inches(0.7), Inches(0.6),
   "→", sz=32, bold=True, color=TEAL)
pic(s, D[6], Inches(5.9), Inches(1.1), Inches(7.0), Inches(5.55))
refline(s, "引用：[9] Self-RAG  [10] RAG Survey — 知识图谱与RAG融合研究基础")

# ═══════════════════════════════════════════════════════════════
# SLIDE 12 — 3-2 任务路由层
# ═══════════════════════════════════════════════════════════════
s = add(content_layout)
header(s, "3-2  任务路由层��计", 12)
rect(s, Inches(0.4), Inches(1.1), Inches(4.8), Inches(0.55), fc=TEAL)
tb(s, Inches(0.5), Inches(1.12), Inches(4.6), Inches(0.5),
   "设计思路", sz=19, bold=True, color=WHITE)
for i,(col,pt,desc) in enumerate([
    (DEEP_BLUE,"意图分类器",
     "识别用户输入属于：事实规则查询 / 空间关系推理 / 约束路线规划 三类任务之一"),
    (TEAL,"动态路由机制",
     "根据分类结果动态分配对应检索模块，避免一刀切的单一路径处理方式"),
    (RED,"核心评估指标",
     "路由准确率（Routing Accuracy）作为核心指标，以单一RAG路径为baseline进行对比实验")]):
    py = Inches(1.75 + i * 1.58)
    rect(s, Inches(0.4), py, Inches(4.8), Inches(1.45), fc=LIGHT_GRAY)
    rect(s, Inches(0.4), py, Inches(0.1), Inches(1.45), fc=col)
    tb(s, Inches(0.6), py+Inches(0.1), Inches(4.5), Inches(0.45),
       pt, sz=16, bold=True, color=col)
    tb(s, Inches(0.6), py+Inches(0.58), Inches(4.5), Inches(0.8),
       desc, sz=13, color=DARK_TEXT)
pic(s, D[4], Inches(5.5), Inches(1.1), Inches(7.4), Inches(5.55))
refline(s, "引用：[8] ReAct  [9] Self-RAG  [6] RAG-Critic")

# ═══════════════════════════════════════════════════════════════
# SLIDE 13 — 3-3 多策略检索+验证层
# ═══════════════════════════════════════════════════════════════
s = add(content_layout)
header(s, "3-3  多��略检索模块与验证层", 13)
pic(s, D[7], Inches(0.4), Inches(1.05), Inches(12.5), Inches(5.9))
refline(s, "引用：[1][2][5][8][11]  — 验证层参考 CRAG [11] 与 Self-RAG [9] 反思机制")

# ═══════════════════════════════════════════════════════════════
# SLIDE 14 — Part 4 过渡
# ═══════════════════════════════════════════════════════════════
section_page(4, "工作计划与预期成果")

# ═══════════════════════════════════════════════════════════════
# SLIDE 15 — 4 工作计划 + 预期成果
# ═══════════════════════════════════════════════════════════════
s = add(content_layout)
header(s, "4  工作计划与预期成果", 15)
rect(s, Inches(0.4), Inches(1.1), Inches(7.2), Inches(0.5), fc=DEEP_BLUE)
tb(s, Inches(0.5), Inches(1.12), Inches(7.0), Inches(0.46),
   "项目进度安排", sz=19, bold=True, color=WHITE)
pic(s, D[5], Inches(0.4), Inches(1.65), Inches(7.2), Inches(4.5))
rect(s, Inches(7.9), Inches(1.1), Inches(5.0), Inches(0.5), fc=TEAL)
tb(s, Inches(8.0), Inches(1.12), Inches(4.8), Inches(0.46),
   "预期成果与经费", sz=19, bold=True, color=WHITE)
for i,(col,ttl,body) in enumerate([
    (TEAL,"发表论文 1 篇",
     "AI或自然语言处理领域\n学术论文（会议/期刊）"),
    (DEEP_BLUE,"软件系统",
     "多策略Agentic RAG系统\n含可视化交互界面"),
    (RED,"经费预算",
     "云计算¥3000 | 印刷¥1000\n版面费¥2000  合计¥6000")]):
    ry = Inches(1.75 + i * 1.65)
    rect(s, Inches(7.9), ry, Inches(5.0), Inches(1.5), fc=LIGHT_GRAY)
    rect(s, Inches(7.9), ry, Inches(0.1), Inches(1.5), fc=col)
    tb(s, Inches(8.1), ry+Inches(0.1), Inches(4.7), Inches(0.45),
       ttl, sz=17, bold=True, color=col)
    tb(s, Inches(8.1), ry+Inches(0.6), Inches(4.7), Inches(0.82),
       body, sz=14, color=DARK_TEXT)

# ═══════════════════════════════════════════════════════════════
# SLIDE 16 — 参考文献
# ═══════════════════════════════════════════════════════════════
s = add(content_layout)
header(s, "参考文献", 16)
refs = [
    "[1]   Lewis P, et al. Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. NeurIPS, 2020.",
    "[2]   Gao L, Ma X, Lin J, Callan J. Precise Zero-Shot Dense Retrieval without Relevance Labels (HyDE). ACL, 2023.",
    "[3]   Es S, James J, Espinosa-Anke L, Schockaert S. RAGAS: Automated Evaluation of RAG. EACL, 2024.",
    "[4]   Saad-Falcon J, et al. ARES: An Automated Evaluation Framework for RAG Systems. NAACL, 2024.",
    "[5]   Liu N F, et al. Lost in the Middle: How Language Models Use Long Contexts. TACL, 2024.",
    "[6]   Dong G, et al. RAG-Critic: Critic-Guided Agentic Workflow for RAG. ACL, 2025.",
    "[7]   Lee D, Jo Y, Park H, Lee M. Shifting from Ranking to Set Selection for RAG (SETR). ACL, 2025.",
    "[8]   Yao S, Zhao J, Yu D, et al. ReAct: Synergizing Reasoning and Acting in Language Models. ICLR, 2023.",
    "[9]   Asai A, Wu Z, Wang Y, Sil A, Hajishirzi H. Self-RAG: Learning to Retrieve, Generate and Critique. ICLR, 2024.",
    "[10]  Gao Y, Xiong Y, Gao X, et al. Retrieval-Augmented Generation for LLMs: A Survey. arXiv, 2024.",
    "[11]  Yan S-Q, Gu J-C, Zhu Y, Ling Z-H. Corrective Retrieval Augmented Generation (CRAG). arXiv, 2024.",
    "[12]  Friel R, Belyi M, Sanyal A. RAGBench: Explainable Benchmark for RAG Systems. arXiv, 2024.",
]
rh = Inches(0.435)
rect(s, Inches(0.4), Inches(1.1), Inches(12.5), rh, fc=DEEP_BLUE)
tb(s, Inches(0.55), Inches(1.12), Inches(12.2), rh-Inches(0.08),
   "文献列表  （共12篇 · RAG及相关技术核心文献）",
   sz=15, bold=True, color=WHITE)
for i,ref in enumerate(refs):
    ry = Inches(1.535 + i * rh)
    rect(s, Inches(0.4), ry, Inches(12.5), rh-Inches(0.02),
         fc=LIGHT_GRAY if i%2==0 else ROW_ALT)
    tb(s, Inches(0.5), ry+Inches(0.03), Inches(12.3), rh-Inches(0.08),
       ref, sz=12.5, color=DARK_TEXT)

# ─── Save ─────────────────────────────────────────────────────
out = "/home/user/claude-dachuang/rag_ppt_v2.pptx"
prs.save(out)
print(f"Saved: {out}")
print(f"Slides: {len(prs.slides)}")
for i,sl in enumerate(prs.slides):
    txts = [s for s in sl.shapes if s.has_text_frame]
    t1 = txts[0].text_frame.paragraphs[0].text[:35] if txts else ""
    print(f"  [{i+1}] {t1}")
